import os
from datetime import date, timezone
from io import BytesIO
from typing import List

from aiohttp import ClientSession
from docx import Document
from fastapi import Depends, HTTPException, Query, status
from openpyxl.workbook import Workbook
from sqlalchemy import and_, desc, exc, func, or_, select, update
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import aliased, joinedload, selectinload
from starlette.responses import StreamingResponse

from src.action_history.models import ActionCode
from src.action_history.schemas import ActionHistoryCreate
from src.action_history.service import action_history_service
from src.common.service import file_service
from src.database import get_db
from src.directions.models import Directions, TransportationType
from src.exceptions import IdNotFound
from src.geography.models import City
from src.notification.models import SmsCode
from src.notification.service import notification_service
from src.orders.models import OrderItems, Orders, OrderStatus, PaymentStatus
from src.orders.schemas import (OrderViewShortSchemas, PaymentType,
                                SendOTPSigning, SignOrderOTP)
from src.orders.service import order_item_service, order_service
from src.shipping.exceptions import ShippingNotFound, ShippingRespondNotFound
from src.shipping.models import (Shipping, ShippingRespond,
                                 ShippingRespondStatus, ShippingStatus,
                                 ShippingWarehouse)
from src.shipping.schemas import (CourierShippingsDetailSchema,
                                  CourierShippingsSchema,
                                  PaginationCourierShippings,
                                  PaginationShipping,
                                  ShippingAddDriverCreateSchema,
                                  ShippingCoordinate, ShippingCreateSchema,
                                  ShippingLoadsSchema, ShippingPathSchema,
                                  ShippingRespondViewSchema,
                                  ShippingViewSchema)
from src.shipping.utils import QueryFilter, enrich_shipping
from src.users.auth import JWTBearer
from src.users.models import Group, OTPSigningCode, OTPType
from src.users.models import Permission
from src.users.models import Permission as UserPermission
from src.users.models import ReviewsDriver, Users, auth_group_permission
from src.users.perms import if_user_has_permissions
from src.users.schemas import (GroupEnum, Permission, UserShippingViewSchemas,
                               UserViewSchemas)
from src.warehouse.models import Warehouse
from src.warehouse.schemas import WarehouseOutShort


class ShippingService:
    async def create(self, payload: ShippingCreateSchema, db: AsyncSession = Depends(get_db)):
        try:
            model = Shipping
            if payload.departure_date is not None:
                payload.departure_date = payload.departure_date.astimezone(
                    timezone.utc).replace(tzinfo=None)
            if payload.arrival_date is not None:
                payload.arrival_date = payload.arrival_date.astimezone(
                    timezone.utc).replace(tzinfo=None)
            set_model = model(status=ShippingStatus.NEW.value,
                              **payload.model_dump(exclude={"warehouses_id"}))
            if payload.warehouses_id:
                for warehouse_id in payload.warehouses_id:
                    warehouse_model = Warehouse
                    warehouse_query = await db.execute(select(warehouse_model).where(warehouse_model.id == warehouse_id))
                    warehouse_data = warehouse_query.scalar_one_or_none()
                    if warehouse_data:
                        new_association = ShippingWarehouse(shipping=set_model, warehouse=warehouse_data)
                        db.add(new_association)
            db.add(set_model)
            await db.commit()
            await db.refresh(set_model)
            return set_model
        except exc.IntegrityError as e:
            await db.rollback()
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail=str(e._message))

    def paginate(self, shippings, page: int, limit: int):
        return shippings.offset((page - 1) * limit).limit(limit)

    def search_shippings(self, shippings, search: str = None):
        if search and not search.isspace():
            arrival_city = aliased(City)
            departure_city = aliased(City)
            shippings = shippings.join(Directions).join(
                arrival_city,
                arrival_city.id == Directions.arrival_city_id).join(
                departure_city,
                departure_city.id == Directions.departure_city_id).where(
                or_(
                    arrival_city.name.ilike(f"%{search}%"),
                    departure_city.name.ilike(f"%{search}%")))
        return shippings

    async def list_shippings(self, is_driver_contract_accepted: bool = Query(None), direction_id: int = Query(None),
                             transportation_type: TransportationType = Query(None),
                             statuses: List[ShippingStatus] = Query(None), db: AsyncSession = Depends(get_db),
                             user: UserViewSchemas = Depends(JWTBearer()), start_date: date = Query(None),
                             end_date: date = Query(None), page: int = 1, limit: int = 10, search: str = Query(None),
                             is_loaded: bool = Query(None)) -> PaginationShipping:
        model = Shipping
        shippings = select(model)
        filtered_shippings = QueryFilter.filter_shippings(
            class_name=model,
            data=shippings,
            status_ids=statuses,
            start_date=start_date,
            end_date=end_date,
            is_driver_contract_accepted=is_driver_contract_accepted,
            transportation_type=transportation_type,
            direction_id=direction_id,
            is_loaded=is_loaded
        )
        if (await if_user_has_permissions(db, user.id, [Permission.ACCEPT_ORDER_TO_WAREHOUSE])):
            condition1 = model.start_warehouse_id == user.warehouse_id, model.status.in_(
                [ShippingStatus.NEW.value, ShippingStatus.WAITING_DRIVER.value, ShippingStatus.IN_TRANSIT.value])
            condition2 = model.end_warehouse_id == user.warehouse_id, model.status.in_(
                [ShippingStatus.IN_TRANSIT.value, ShippingStatus.FINISHED.value])
            filtered_shippings = filtered_shippings.where(
                or_(and_(*condition1), and_(*condition2)))
        filtered_shippings = self.search_shippings(
            filtered_shippings, search).order_by(desc(model.created_at))
        paginated_shippings = self.paginate(filtered_shippings, page, limit)
        query = await db.execute(paginated_shippings)
        shippings = query.scalars().all()
        shippings = [await enrich_shipping(shipping, db=db, user=user) for shipping in shippings]
        shippings = [
            shipping for shipping in shippings if shipping.is_canceled == False]
        total = await db.execute(filtered_shippings)
        total = len(total.scalars().all())
        pages_number = total // limit
        if total % limit != 0:
            pages_number += 1
        return PaginationShipping(
            page=page,
            pages_number=pages_number,
            total=total,
            limit=limit,
            data=shippings
        )

    async def filter_my_shippings(self, shippings, responds, db: AsyncSession = Depends(get_db), user: Users = Depends(JWTBearer())) -> PaginationShipping:
        responds = responds.where(ShippingRespond.driver_id == user.id)
        responds = await db.execute(responds)
        responds = responds.scalars().all()
        shipping_ids = [item.shipping_id for item in responds]
        shippings = shippings.where(Shipping.id.in_(shipping_ids))
        return shippings

    async def my_shippings(self, is_driver_contract_accepted: bool = Query(None), respond_status: list[ShippingRespondStatus] = Query(None), direction_id: int = Query(None), transportation_type: TransportationType = Query(None), statuses: List[ShippingStatus] = Query(None), db: AsyncSession = Depends(get_db), user: Users = Depends(JWTBearer()), start_date: date = Query(None), end_date: date = Query(None), page: int = 1, limit: int = 10, search: str = Query(None)) -> PaginationShipping:
        shippings = select(Shipping)
        responds = select(ShippingRespond)
        if respond_status:
            responds = responds.where(
                ShippingRespond.respond_status.in_(respond_status))
        filtered_shippings = await self.filter_my_shippings(
            shippings=shippings,
            responds=responds,
            db=db,
            user=user
        )
        filtered_shippings = QueryFilter.filter_shippings(
            class_name=Shipping,
            data=filtered_shippings,
            status_ids=statuses,
            start_date=start_date,
            end_date=end_date,
            is_driver_contract_accepted=is_driver_contract_accepted,
            transportation_type=transportation_type,
            direction_id=direction_id
        )
        filtered_shippings = self.search_shippings(filtered_shippings, search)
        paginated_shippings = self.paginate(filtered_shippings, page, limit)
        query = await db.execute(paginated_shippings)
        shippings = query.scalars().all()
        shippings = [await enrich_shipping(shipping, db=db, user=user) for shipping in shippings]
        total = await db.execute(filtered_shippings)
        total = len(total.scalars().all())
        pages_number = total // limit
        if total % limit != 0:
            pages_number += 1
        return PaginationShipping(
            page=page,
            pages_number=pages_number,
            total=total,
            limit=limit,
            data=shippings
        )

    async def get_by_id(self, id: int = None, db: AsyncSession = Depends(get_db), user: Users = Depends(JWTBearer())) -> ShippingViewSchema:
        model = Shipping
        query = await db.execute(select(model).options(selectinload(Shipping.orders).selectinload(Orders.direction)).where(model.id == id))
        shipping = query.scalar_one_or_none()
        return await enrich_shipping(shipping, db=db, user=user)

    async def path(self, id: int, payload: ShippingPathSchema, db: AsyncSession = Depends(get_db)) -> ShippingViewSchema:
        if not id:
            raise IdNotFound()
        model = Shipping
        query = await db.execute(select(model).where(model.id == id))
        data = query.scalar_one_or_none()
        if not data:
            raise ShippingNotFound()
        try:
            stmt = (
                update(Shipping)
                .where(Shipping.id == id)
                .values(**payload.model_dump(exclude_unset=True))
            )
            await db.execute(stmt)
            await db.commit()
            await db.refresh(data)
            # TODO error of MissingGreenlet, temporary resolve of issue. original
            # error: greenlet_spawn has not been called; can't call await_only() here.
            # Was IO attempted in an unexpected place? (Background on this error at:
            # https://sqlalche.me/e/20/xd2s)
            warehouse_shipping_query = select(ShippingWarehouse).options(
                selectinload(ShippingWarehouse.warehouse)).where(
                ShippingWarehouse.shipping_id == id
            )
            warehouse_shipping_select = await db.execute(warehouse_shipping_query)
            warehouse_shipping_data = warehouse_shipping_select.scalars().all()
            warehouse_shipping_out = [
                WarehouseOutShort(id=ws_data.warehouse.id, name=ws_data.warehouse.name,
                                  address=ws_data.warehouse.address)
                for ws_data in warehouse_shipping_data if ws_data.warehouse
            ]
            return ShippingViewSchema(**data.__dict__, warehouses=warehouse_shipping_out)
        except exc.IntegrityError as e:
            await db.rollback()
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail=str(e._message))

    async def delete(self, id: int, db: AsyncSession = Depends(get_db)) -> None:
        if not id:
            raise IdNotFound()
        model = Shipping
        query = await db.execute(select(model).where(model.id == id))
        data = query.scalar_one_or_none()
        if not data:
            raise ShippingNotFound()
        await db.delete(data)
        await db.commit()

    async def add_driver(self, id: int, payload: ShippingAddDriverCreateSchema, db: AsyncSession = Depends(get_db)) -> ShippingViewSchema:
        model = Shipping
        query = await db.execute(select(model).where(model.id == id))
        data = query.scalar_one_or_none()
        if len(data.drivers) > 0:
            data.drivers = []
        if data is None:
            raise ShippingNotFound()
        for user_id in payload.drivers:
            user_model = Users
            user_query = await db.execute(select(user_model).where(user_model.id == user_id))
            user_data = user_query.scalar_one_or_none()
            if user_data is None:
                continue
            else:
                data.drivers.append(user_data)
        await db.commit()
        await db.refresh(data)
        return data

    def generate_otp_code(self) -> str:
        import random
        return str(random.randint(100000, 999999))

    async def generate_driver_contract(self, driver: Users, otp_code: str, order: Shipping, db: AsyncSession = Depends(get_db)):
        doc_path = f'ip_agreement_for_driver.docx'
        doc = Document(doc_path)
        sended_date = date.today().strftime('%m/%d/%Y')
        replacements = {
            '{{driver_fio}}': driver.fl_name if driver else '',
            '{{sended_date}}': sended_date,
            '{{driver_phone}}': driver.phone if driver else '',
            '{{otp_code}}': otp_code
        }
        for paragraph in doc.paragraphs:
            for key, value in replacements.items():
                if key in paragraph.text:
                    paragraph.text = paragraph.text.replace(key, value)

        new_doc_path = f'ip_agreement_for_driver_{otp_code}.docx'
        doc.save(new_doc_path)
        new_pdf_path = f'ip_agreement_for_driver_{otp_code}.pdf'
        file_stream = await order_service.convert_docx_to_pdf(new_doc_path, new_pdf_path)
        file_s3 = await file_service.upload_file(new_pdf_path, file_stream)
        os.remove(new_doc_path)
        order.driver_contract_url = file_s3['file_path']
        await db.commit()

    async def get_driver_contract(self, id: int, user: UserViewSchemas, db: AsyncSession = Depends(get_db)) -> str:
        shipping = await db.get_one(Shipping, id)
        if shipping is None:
            raise ShippingNotFound()
        file_response = await file_service.download_file(f'{shipping.driver_contract_url}')
        return file_response

    async def send_otp_code_for_signing(self, id: int, user: UserViewSchemas, db: AsyncSession = Depends(get_db)) -> SendOTPSigning:
        query = await db.execute(select(Shipping).where(Shipping.id == id))
        order: Shipping = query.scalar_one_or_none()
        if order is None:
            raise HTTPException(status_code=404, detail="Shipping not found")
        otp_code = self.generate_otp_code()
        driver = await db.execute(select(Users).where(Users.id == order.driver_id))
        driver = driver.scalar_one_or_none()
        notification_service.send_sms(
            phone=driver.phone,
            sms_code=SmsCode.DRIVER_CONTRACT,
            code=otp_code)
        otp_signing_code = OTPSigningCode(
            user_id=user.id,
            code=otp_code,
            phone=user.phone,
            shipping_id=id,
            otp_type=OTPType.DRIVER_CONTRACT.value)
        db.add(otp_signing_code)
        await db.commit()
        await self.generate_driver_contract(driver, otp_code, order, db)
        return SendOTPSigning(phone=driver.phone, code=otp_code)

    async def accept_driver_contract(self, id: int, payload: SignOrderOTP, user: UserViewSchemas, db: AsyncSession = Depends(get_db)):
        query = await db.execute(select(Shipping).where(Shipping.id == id))
        shipping = query.scalar_one_or_none()
        if shipping is None:
            raise HTTPException(status_code=404, detail="OTP not found")
        otp_signing_code = await db.execute(select(OTPSigningCode).where(OTPSigningCode.code == payload.code))
        otp_signing_code = otp_signing_code.scalar_one_or_none()
        if otp_signing_code is None:
            raise HTTPException(status_code=404, detail="OTP code not found")
        await db.execute(update(Shipping).where(Shipping.id == id).values(is_driver_contract_accepted=True))
        await db.execute(update(OTPSigningCode).where(OTPSigningCode.id == otp_signing_code.id).values(is_used=True))
        await db.commit()
        return otp_signing_code

    async def create_shipping_responds(self, shipping_id, user: Users = Depends(JWTBearer()), db: AsyncSession = Depends(get_db)) -> ShippingRespondViewSchema:
        shipping_model = Shipping
        shipping_query = await db.execute(select(shipping_model).where(shipping_model.id == shipping_id))
        shipping_data = shipping_query.scalar_one_or_none()
        if shipping_data is None:
            raise ShippingNotFound()
        model = ShippingRespond
        shipping_respond = select(model).where(
            model.shipping_id == shipping_id, model.driver_id == user.id)
        shipping_respond = await db.execute(shipping_respond)
        shipping_respond = shipping_respond.scalar_one_or_none()
        if shipping_respond:
            raise HTTPException(
                status_code=400, detail="You already responded")
        set_model = model(shipping_id=shipping_id, driver_id=user.id,
                          respond_status=ShippingRespondStatus.RESPONDED.value)
        db.add(set_model)
        await db.commit()
        await db.refresh(set_model)
        reviews_number = select(
            func.count(
                ReviewsDriver.id)).where(
            ReviewsDriver.driver_id == set_model.driver_id)
        reviews_number = await db.scalar(reviews_number)
        driver = UserShippingViewSchemas(
            **set_model.driver.__dict__,
            reviews_number=reviews_number if reviews_number else 0)
        shipping = ShippingViewSchema(**set_model.shipping.__dict__)
        return ShippingRespondViewSchema(
            id=set_model.id,
            driver=driver,
            shipping=shipping,
            respond_status=ShippingRespondStatus.RESPONDED.value)

    async def get_shipping_responds_by_shipping(self, shipping_id: int, status: ShippingRespondStatus = Query(None), db: AsyncSession = Depends(get_db)) -> List[ShippingRespondViewSchema]:
        model = ShippingRespond
        filtered_model = QueryFilter.filter_shippings(
            class_name=model, data=select(model), status=status).where(
            model.shipping_id == shipping_id)
        query = await db.execute(filtered_model)
        data = query.scalars().all()
        result = []
        for item in data:
            reviews_number = select(
                func.count(
                    ReviewsDriver.id)).where(
                ReviewsDriver.driver_id == item.driver_id)
            reviews_number = await db.scalar(reviews_number)
            driver = UserShippingViewSchemas(
                **item.driver.__dict__,
                reviews_number=reviews_number if reviews_number else 0)
            # TODO error of MissingGreenlet, temporary resolve of issue. original
            # error: greenlet_spawn has not been called; can't call await_only() here.
            # Was IO attempted in an unexpected place? (Background on this error at:
            # https://sqlalche.me/e/20/xd2s)
            if item.shipping.orders:
                orders_data = []
                for order in item.shipping.orders:
                    order_dict = order.__dict__
                    order_dict.pop('direction_id', None)
                    order_direction = order.direction
                    order_data = OrderViewShortSchemas(
                        **order_dict,
                        direction_id=order_direction
                    )
                    orders_data.append(order_data)
                item_shipping_dict = item.shipping.__dict__
                item_shipping_dict.pop('orders')
                shipping = ShippingViewSchema(**item_shipping_dict, orders=orders_data)
            else:
                shipping = ShippingViewSchema(**item.shipping.__dict__)
            result.append(
                ShippingRespondViewSchema(
                    id=item.id,
                    driver=driver,
                    shipping=shipping))
        return result

    async def accept(self, respond_id: int, user: UserViewSchemas, db: AsyncSession = Depends(get_db)):
        try:
            model = ShippingRespond
            query = await db.get_one(model, respond_id)
            if query is None:
                raise ShippingRespondNotFound()
            query.respond_status = ShippingRespondStatus.CONFIRMED.value
            model_shipping = Shipping
            query_shipping = await db.get_one(model_shipping, query.shipping_id)
            if query_shipping is None:
                raise ShippingNotFound()
            query_shipping.driver_id = query.driver_id
            query_shipping.status = ShippingStatus.WAITING_DRIVER.value
            await db.execute(
                update(model)
                .where(model.id != respond_id)
                .values(respond_status=ShippingRespondStatus.CANCEL.value)
            )
            await db.commit()
            await db.refresh(query)
            return {"detail": "Водитель нанят"}
        except exc.NoResultFound:
            await db.rollback()

    async def cancel(self, respond_id: int, user: UserViewSchemas, db: AsyncSession = Depends(get_db)):
        try:
            model = ShippingRespond
            query = await db.get_one(model, respond_id)
            if query is None:
                raise ShippingRespondNotFound()
            query.respond_status = ShippingRespondStatus.CANCEL.value
            model_shipping = Shipping
            query_shipping = await db.get_one(model_shipping, query.shipping_id)
            if query_shipping is None:
                raise ShippingNotFound()
            query_shipping.driver_id = query.driver_id
            query_shipping.status = ShippingStatus.NEW.value
            await db.commit()
            await db.refresh(query)
            return query
        except exc.NoResultFound:
            await db.rollback()

    async def shipping_loads(self, payload: ShippingLoadsSchema, shipping_id: int, user: UserViewSchemas, db: AsyncSession = Depends(get_db)):
        model = Shipping
        query = await db.execute(select(model).where(model.id == shipping_id).options(selectinload(Shipping.warehouses)))
        shipping = query.scalar_one_or_none()
        if shipping is None:
            raise ShippingNotFound()
        associated_warehouse_ids = [warehouse.warehouse_id for warehouse in shipping.warehouses]
        all_orders_query = await db.execute(
            select(Orders).options(
                selectinload(Orders.direction)).join(OrderItems).filter(OrderItems.id.in_(payload.orders_items_id)))
        all_orders = all_orders_query.scalars().all()
        total_weight = 0
        total_volume = 0
        for order in all_orders:
            total_weight += order.total_weight
            total_volume += order.total_volume
            if shipping.shipping_type == TransportationType.ROAD.value:
                if order.destination_warehouse_id != shipping.end_warehouse_id and order.destination_warehouse_id not in associated_warehouse_ids:
                    raise HTTPException(
                        status_code=400,
                        detail=f"Номер заказа: {order.id}. Склад по месту назначения: {order.destination_warehouse_id} не совпадает с склад-остановкой перевозки")
            elif shipping.shipping_type == TransportationType.AIR.value or shipping.shipping_type == TransportationType.RAIL.value:
                if order.direction.arrival_city_id != shipping.direction.arrival_city_id:
                    raise HTTPException(
                        status_code=400,
                        detail=f"Номер заказа: {order.id}. Город прибытия заказа: {order.direction.arrival_city_id} не совпадает с Городом прибытия перевозки")
        if shipping.shipping_type == TransportationType.RAIL.value or shipping.shipping_type == TransportationType.AIR.value:
            shipping.cargo_weight = total_weight
            shipping.cargo_volume = total_volume
        shipping.is_loaded = True
        for item_id in payload.orders_items_id:
            order_item_model = OrderItems
            order_item_query = await db.execute(select(order_item_model).where(order_item_model.id == item_id))
            order_item_data = order_item_query.scalar_one_or_none()
            order_items = await db.execute(select(OrderItems).where(OrderItems.order_id == order_item_data.order_id))
            order_items = order_items.scalars().all()
            if order_item_data:
                order_item_data.is_loaded = True
                await order_item_service.set_as_in_transit(item_id, user, db)
                shipping.order_items.append(order_item_data)
                order_items = await db.execute(
                    select(OrderItems).where(OrderItems.order_id == order_item_data.order_id))
                order_items = order_items.scalars().all()
                in_transit_check = all(
                    item.status == OrderStatus.IN_TRANSIT.value for item in order_items)
                order_query = await db.execute(select(Orders).where(Orders.id == order_item_data.order_id))
                order = order_query.scalar_one_or_none()
                if order:
                    shipping.orders.append(order)
                    if in_transit_check:
                        await order_service.set_as_in_transit(order.id, user, db)
                        if (shipping.shipping_type in [TransportationType.RAIL.value, TransportationType.AIR.value]):
                            await db.execute(update(Shipping).where(Shipping.id == shipping_id).values(status=ShippingStatus.IN_TRANSIT.value))
                    else:
                        await order_service.set_as_partial_in_transit(order.id, user, db)
                    await db.commit()
                    await db.refresh(order)
                else:
                    await order_item_service.set_as_in_transit(item_id, user, db)
                    shipping.order_items.append(order_item_data)
                    order = await db.execute(select(Orders).where(Orders.id == order_item_data.order_id))
                    order = order.scalar_one_or_none()
                    if order:
                        shipping.orders.append(order)
                        await order_service.set_as_in_transit(order.id, user, db)
                        if (shipping.shipping_type in [TransportationType.RAIL.value, TransportationType.AIR.value]):
                            await db.execute(update(Shipping).where(Shipping.id == shipping_id).values(status=ShippingStatus.IN_TRANSIT.value))
        await db.commit()
        await db.refresh(shipping)
        return {"detail": "Погрузка прошла успешно"}

    async def status_in_transit(self, shipping_id: int, db: AsyncSession = Depends(get_db)) -> ShippingViewSchema:
        model = Shipping
        query = await db.execute(select(model).where(model.id == shipping_id))
        data = query.scalar_one_or_none()
        if data is None:
            raise ShippingNotFound()
        if not data.is_loaded:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST, detail="не загружен, не можете начать поездку")
        data.status = ShippingStatus.IN_TRANSIT.value
        await db.commit()
        await db.refresh(data)
        # TODO error of MissingGreenlet, temporary resolve of issue. original
        # error: greenlet_spawn has not been called; can't call await_only() here.
        # Was IO attempted in an unexpected place? (Background on this error at:
        # https://sqlalche.me/e/20/xd2s)
        orders_data = []
        for order in data.orders:
            order_dict = order.__dict__
            order_dict.pop('direction_id', None)
            order_direction = order.direction
            order_data = OrderViewShortSchemas(
                **order_dict,
                direction_id=order_direction
            )
            orders_data.append(order_data)
        shipping_dict = data.__dict__
        shipping_dict.pop('orders')
        return ShippingViewSchema(
            **shipping_dict,
            orders=orders_data
        )

    async def status_finished(self, shipping_id: int, db: AsyncSession = Depends(get_db)) -> None:
        model = Shipping
        query = await db.execute(select(model).where(model.id == shipping_id))
        data = query.scalar_one_or_none()
        if data is None:
            raise ShippingNotFound()
        data.status = ShippingStatus.FINISHED.value
        shipping_respond = select(ShippingRespond).where(
            ShippingRespond.shipping_id == shipping_id,
            ShippingRespond.respond_status == ShippingRespondStatus.CONFIRMED.value,
            ShippingRespond.driver_id == data.driver_id)
        shipping_respond = await db.execute(shipping_respond)
        shipping_respond = shipping_respond.scalar_one_or_none()
        if shipping_respond:
            shipping_respond.respond_status = ShippingRespondStatus.FINISHED.value
        await db.commit()
        await db.refresh(data)
        return data

    async def coordinates(self, payload: ShippingCoordinate, shipping_id: int, db: AsyncSession = Depends(get_db)) -> ShippingViewSchema:
        model = Shipping
        query = await db.execute(select(model).where(model.id == shipping_id))
        data = query.scalar_one_or_none()
        if data is None:
            raise ShippingNotFound()
        data.longitude = payload.longitude
        data.latitude = payload.latitude
        await db.commit()
        await db.refresh(data)
        # TODO error of MissingGreenlet, temporary resolve of issue. original
        # error: {'type': 'get_attribute_error', 'loc': ('response',
        # 'warehouses'), 'msg': "Error extracting attribute: MissingGreenlet:
        # greenlet_spawn has not been called; can't call await_only() here. Was IO
        # attempted in an unexpected place? (Background on this error at:
        # https://sqlalche.me/e/20/xd2s)"
        warehouse_shipping_query = (
            select(ShippingWarehouse)
            .join(Warehouse)
            .options(selectinload(ShippingWarehouse.warehouse))
            .where(ShippingWarehouse.shipping_id == Shipping.id)
            .distinct(ShippingWarehouse.warehouse_id)
            .order_by(ShippingWarehouse.warehouse_id)
        )
        warehouse_shipping_select = await db.execute(warehouse_shipping_query)
        warehouse_shipping_data = warehouse_shipping_select.scalars().all()
        warehouse_shipping_out = [
            WarehouseOutShort(id=ws_data.warehouse.id, name=ws_data.warehouse.name, address=ws_data.warehouse.address)
            for ws_data in warehouse_shipping_data if ws_data.warehouse
        ]
        return ShippingViewSchema(**data.__dict__, warehouses=warehouse_shipping_out)

    async def generate_excel_file(self, id: int, db: AsyncSession = Depends(get_db),
                                  user: UserViewSchemas = Depends(JWTBearer())):
        async with db.begin():
            query = select(Orders).join(
                Orders.shippings
            ).options(
                selectinload(Orders.warehouse),
                selectinload(Orders.payment),
                selectinload(Orders.direction)
            ).where(
                Shipping.id == id).order_by(Orders.created_at.asc())
            filtered_result = await db.execute(query)
            filtered_orders = filtered_result.scalars().all()
            order_items_count_query = select(
                OrderItems.order_id,
                func.count().label('items_count')
            ).group_by(OrderItems.order_id)

            result = await db.execute(order_items_count_query)
            order_items_count = result.all()
        order_items_count_dict = {order_id: count for order_id, count in order_items_count}
        wb = Workbook()
        ws = wb.active
        ws.title = "Заказы"

        columns = [
            "ID",
            "Дата",
            "Откуда",
            "Отправитель",
            "Контакт отправителя",
            "Количество",
            "Вес",
            "Объем",
            "Сумма",
            "Упаковка",
            "Доставка",
            "Забор",
            "Куда",
            "Получатель",
            "Контакты получателя",
            "Тип отправки",
            "Тип оплаты",
            "Местонахождение",
            "Оплачено",
            "Сотрудник",
            "Характер груза"]
        ws.append(columns)
        for order in filtered_orders:
            payment_status = 'ДА' if order.payment and order.payment.payment_status == PaymentStatus.PAID.value else 'НЕТ'
            payment_type = None
            transportation_type = None
            if order.payment:
                if order.payment.payment_type == PaymentType.CASH.value:
                    payment_type = 'Наличные'
                if order.payment.payment_type == PaymentType.ONLINE.value:
                    payment_type = 'Онлайн'
            if order.direction:
                if order.direction.transportation_type == TransportationType.AIR.value:
                    transportation_type = 'Самолет'
                if order.direction.transportation_type == TransportationType.RAIL.value:
                    transportation_type = 'ЖД'
                if order.direction.transportation_type == TransportationType.ROAD.value:
                    transportation_type = 'Фура'
            ws.append([
                order.id,
                order.created_at,
                order.direction.departure_city.name if order.direction else None,
                order.sender_fio,
                order.sender_phone,
                order_items_count_dict.get(order.id, 0),
                order.total_weight,
                order.total_volume,
                order.payment.amount if order.payment else 0,
                None,  # Упаковка
                None,  # Доставка
                None,  # Забор
                order.direction.arrival_city.name if order.direction else None,
                order.receiver_fio,
                order.receiver_phone,
                transportation_type,
                payment_type,
                f'{order.warehouse.name}, {order.warehouse.address}' if order.warehouse else None,
                payment_status,
                None,  # Сотрудник,
                order.description

            ])
        excel_file = BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)
        return excel_file

    async def get_courier(self, db: AsyncSession = Depends(get_db), page: int = 1, limit: int = 10, search: str = Query(None)) -> PaginationCourierShippings:
        query = select(Users).join(
            auth_group_permission, Users.group_id == auth_group_permission.c.group_id
        ).join(
            UserPermission, UserPermission.codename == auth_group_permission.c.codename
        ).where(
            UserPermission.codename == Permission.DELIVER_ORDER.value
        )

        if search:
            query = query.filter(or_(
                Users.first_name.ilike(f"%{search}%"),
                Users.last_name.ilike(f"%{search}%")
            ))
        total_query = await db.execute(query)  # This is to calculate total entries matching the filter
        total = total_query.scalars().all()
        total_count = len(total)

        query = query.offset((page - 1) * limit).limit(limit)
        result = await db.execute(query)
        users = result.scalars().all()

        courier_data = [CourierShippingsSchema(
            id=user.id,
            name=user.fl_name,
            order_count=len((await db.execute(
                select(Orders)
                .where(Orders.courier == user.id)
            )).scalars().all())
        ) for user in users]

        pages_number = total_count // limit
        if total_count % limit != 0:
            pages_number += 1

        return PaginationCourierShippings(
            page=page,
            limit=limit,
            total=total_count,
            pages_number=pages_number,
            data=courier_data
        )

    async def get_courier_detail(self, id: int, db: AsyncSession = Depends(get_db)) -> CourierShippingsDetailSchema:
        result = await db.execute(
            select(Users).join(
                auth_group_permission, Users.group_id == auth_group_permission.c.group_id
            ).join(UserPermission, UserPermission.codename == auth_group_permission.c.codename).where(Users.id == id,
                                                                                                      UserPermission.codename == Permission.DELIVER_ORDER.value)
        )
        user = result.scalars().first()
        if not user:
            raise HTTPException(status_code=404, detail="Courier not found or does not have delivery permissions.")

        orders_result = await db.execute(
            select(Orders).options(joinedload(Orders.direction))
            .where(Orders.courier == user.id).order_by(Orders.created_at.asc())
        )
        orders = orders_result.scalars().all()

        orders_data = [
            OrderViewShortSchemas(
                **order.__dict__
            ) for order in orders
        ]

        return CourierShippingsDetailSchema(
            id=user.id,
            name=user.fl_name,
            order_count=len(orders_data),
            orders=orders_data
        )

    async def generate_courier_excel(self, id: int, db: AsyncSession = Depends(get_db),
                                     user: UserViewSchemas = Depends(JWTBearer())):
        async with db.begin():
            query = select(Orders).options(
                selectinload(Orders.warehouse),
                selectinload(Orders.payment),
                selectinload(Orders.direction)
            ).where(
                Orders.courier == id).order_by(Orders.created_at.asc())
            filtered_result = await db.execute(query)
            filtered_orders = filtered_result.scalars().all()
            order_items_count_query = select(
                OrderItems.order_id,
                func.count().label('items_count')
            ).group_by(OrderItems.order_id)

            result = await db.execute(order_items_count_query)
            order_items_count = result.all()
        order_items_count_dict = {order_id: count for order_id, count in order_items_count}
        wb = Workbook()
        ws = wb.active
        ws.title = "Заказы"

        columns = [
            "ID",
            "Дата",
            "Откуда",
            "Отправитель",
            "Контакт отправителя",
            "Количество",
            "Вес",
            "Объем",
            "Сумма",
            "Упаковка",
            "Доставка",
            "Забор",
            "Куда",
            "Получатель",
            "Контакты получателя",
            "Тип отправки",
            "Тип оплаты",
            "Местонахождение",
            "Оплачено",
            "Сотрудник",
            "Характер груза"]
        ws.append(columns)
        for order in filtered_orders:
            payment_status = 'ДА' if order.payment and order.payment.payment_status == PaymentStatus.PAID.value else 'НЕТ'
            payment_type = None
            transportation_type = None
            if order.payment:
                if order.payment.payment_type == PaymentType.CASH.value:
                    payment_type = 'Наличные'
                if order.payment.payment_type == PaymentType.ONLINE.value:
                    payment_type = 'Онлайн'
            if order.direction:
                if order.direction.transportation_type == TransportationType.AIR.value:
                    transportation_type = 'Самолет'
                if order.direction.transportation_type == TransportationType.RAIL.value:
                    transportation_type = 'ЖД'
                if order.direction.transportation_type == TransportationType.ROAD.value:
                    transportation_type = 'Фура'
            ws.append([
                order.id,
                order.created_at,
                order.direction.departure_city.name if order.direction else None,
                order.sender_fio,
                order.sender_phone,
                order_items_count_dict.get(order.id, 0),
                order.total_weight,
                order.total_volume,
                order.payment.amount if order.payment else 0,
                None,  # Упаковка
                None,  # Доставка
                None,  # Забор
                order.direction.arrival_city.name if order.direction else None,
                order.receiver_fio,
                order.receiver_phone,
                transportation_type,
                payment_type,
                f'{order.warehouse.name}, {order.warehouse.address}' if order.warehouse else None,
                payment_status,
                None,  # Сотрудник,
                order.description

            ])
        excel_file = BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)
        return excel_file


shipping_service = ShippingService()
